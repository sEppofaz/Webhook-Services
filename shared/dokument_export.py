"""Geschriebene Dokumente (Formular) als PDF oder Word (ADR-029, v1.78) – entsteht bei jedem Abruf neu.

Protokoll: Kopf (Verein, Sitzung, Datum, Ort, Zeit, Leitung, Protokoll, Anwesend, Entschuldigt), Tagesordnungspunkte
mit Beschluss und Abstimmung, Freitext. Satzung/Sonstiges: Titel, Datum, Freitext. Schrift wie `shared/export.py`.
"""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

from shared.export import _SCHRIFTEN

FORMATE = {"pdf": ("PDF", "application/pdf"),
           "docx": ("Word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}


def datum_lang(iso: str) -> str:
    try:
        return datetime.strptime(iso, "%Y-%m-%d").strftime("%d.%m.%Y")
    except (TypeError, ValueError):
        return ""


def abstimmung(top: dict) -> str:
    teile = [f"{n} {top[k]}" for k, n in (("ja", "Ja"), ("nein", "Nein"), ("enthaltung", "Enthaltung")) if str(top.get(k, "")).strip()]
    return " · ".join(teile)


def kopfzeilen(dok: dict, verein_name: str) -> list[tuple[str, str]]:
    i = dok.get("inhalt") or {}
    zeit = " – ".join(x for x in (i.get("beginn", ""), i.get("ende", "")) if x)
    zeilen = [("Verein", verein_name)]
    if dok["kategorie"] == "protokoll":
        zeilen += [("Sitzung", i.get("sitzungsart", "")), ("Datum", datum_lang(dok.get("datum", ""))),
                   ("Ort", i.get("ort", "")), ("Uhrzeit", f"{zeit} Uhr" if zeit else ""),
                   ("Sitzungsleitung", i.get("leitung", "")), ("Protokoll", i.get("protokoll", "")),
                   ("Anwesend", i.get("anwesende", "")), ("Entschuldigt", i.get("entschuldigt", ""))]
    else:
        zeilen.append(("Datum", datum_lang(dok.get("datum", ""))))
    return [(k, v) for k, v in zeilen if v]


def _pdf(dok: dict, verein_name: str) -> bytes:
    from fpdf import FPDF
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    schrift = "Helvetica"
    for normal, fett in _SCHRIFTEN:
        if Path(normal).exists() and Path(fett).exists():
            pdf.add_font("Text", "", normal)
            pdf.add_font("Text", "B", fett)
            schrift = "Text"
            break

    def txt(s) -> str:
        s = str(s or "")
        if schrift != "Helvetica":
            return s
        s = s.replace("–", "-").replace("„", '"').replace("“", '"').replace("…", "...").replace("·", "-")
        return s.encode("latin-1", "replace").decode("latin-1")

    def absatz(s: str, groesse: int = 10, fett: bool = False, hoehe: float = 5.5) -> None:
        pdf.set_font(schrift, "B" if fett else "", groesse)
        pdf.multi_cell(0, hoehe, txt(s), new_x="LMARGIN", new_y="NEXT")

    pdf.set_title(dok["titel"])
    pdf.set_creation_date(datetime.fromisoformat(dok["geaendert_am"]))   # gleicher Stand = gleiche Datei
    pdf.add_page()
    absatz(dok["titel"], 16, True, 8)
    pdf.ln(2)
    for k, v in kopfzeilen(dok, verein_name):
        pdf.set_font(schrift, "B", 10)
        pdf.cell(36, 5.5, txt(k))
        pdf.set_font(schrift, "", 10)
        pdf.multi_cell(0, 5.5, txt(v), new_x="LMARGIN", new_y="NEXT")
    i = dok.get("inhalt") or {}
    for n, top in enumerate(i.get("tops") or [], 1):
        pdf.ln(3)
        absatz(f"TOP {n}: {top.get('titel', '')}", 12, True, 6.5)
        if top.get("text"):
            absatz(top["text"])
        if top.get("beschluss"):
            absatz("Beschluss: " + top["beschluss"], fett=True)
        if abstimmung(top):
            absatz("Abstimmung: " + abstimmung(top))
    if i.get("text"):
        pdf.ln(3)
        if dok["kategorie"] == "protokoll":
            absatz("Sonstiges", 12, True, 6.5)
        absatz(i["text"])
    if dok["kategorie"] == "protokoll":
        pdf.ln(14)
        pdf.set_font(schrift, "", 9)
        breite = (pdf.w - pdf.l_margin - pdf.r_margin - 10) / 2
        y = pdf.get_y()
        pdf.line(pdf.l_margin, y, pdf.l_margin + breite, y)
        pdf.line(pdf.l_margin + breite + 10, y, pdf.w - pdf.r_margin, y)
        pdf.cell(breite + 10, 5, txt("Sitzungsleitung"))
        pdf.cell(0, 5, txt("Protokoll"), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _docx(dok: dict, verein_name: str) -> bytes:
    from docx import Document
    from docx.shared import Pt
    d = Document()
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(dok["titel"], level=1)
    tab = d.add_table(rows=0, cols=2)
    for k, v in kopfzeilen(dok, verein_name):
        zeile = tab.add_row().cells
        zeile[0].text = k
        zeile[0].paragraphs[0].runs[0].bold = True
        zeile[1].text = v
    i = dok.get("inhalt") or {}
    for n, top in enumerate(i.get("tops") or [], 1):
        d.add_heading(f"TOP {n}: {top.get('titel', '')}", level=2)
        if top.get("text"):
            d.add_paragraph(top["text"])
        if top.get("beschluss"):
            p = d.add_paragraph()
            p.add_run("Beschluss: ").bold = True
            p.add_run(top["beschluss"])
        if abstimmung(top):
            d.add_paragraph("Abstimmung: " + abstimmung(top))
    if i.get("text"):
        if dok["kategorie"] == "protokoll":
            d.add_heading("Sonstiges", level=2)
        d.add_paragraph(i["text"])
    if dok["kategorie"] == "protokoll":
        d.add_paragraph("\n\n______________________________          ______________________________")
        d.add_paragraph("Sitzungsleitung                                         Protokoll")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def exportiere(fmt: str, dok: dict, verein_name: str) -> bytes:
    return {"pdf": _pdf, "docx": _docx}[fmt](dok, verein_name)
