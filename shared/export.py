"""Terminlisten exportieren: PDF, Word, Excel, LibreOffice (Text/Tabelle), Kalenderdatei.

Prototyp auf Branch `jahresplanung` (Josef 2026-10-05): Nach dem Planungstreffen lädt jeder
Verein seine vereinbarten Termine herunter, die Organisatorin das Jahresprogramm der Gemeinde.

Alle Formate bekommen dieselben Zeilen: Dicts mit `datum` (YYYY-MM-DD), `uhrzeit`, `uhrzeit_bis`,
`bezeichnung`, `ort`, `verein_name`. `exportiere(fmt, titel, zeilen, mit_verein)` → (bytes, mimetype,
dateiendung).

Bibliotheken: openpyxl und python-docx liegen schon im Server-venv; fpdf2 und odfpy wären neu
(beide reines Python). Werden erst beim Aufruf importiert – fehlt eine, scheitert nur ihr Format.

**Später zusammenführen:** `_ics_escape`/`_ics_zeiten` sind Kopien aus `services/kalender/routes.py`.
Die Routen-Datei zieht beim Import `kalender_core` (braucht Secrets) – deshalb hier eigenständig.
Bei der Integration beide nach `shared/` ziehen und hier wie dort nutzen.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from shared.wiederholung import MONATE

WT_KURZ = ("Mo", "Di", "Mi", "Do", "Fr", "Sa", "So")

FORMATE = {
    "pdf":  ("PDF", "application/pdf"),
    "docx": ("Word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "xlsx": ("Excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "odt":  ("LibreOffice Text", "application/vnd.oasis.opendocument.text"),
    "ods":  ("LibreOffice Tabelle", "application/vnd.oasis.opendocument.spreadsheet"),
    "ics":  ("Kalenderdatei", "text/calendar; charset=utf-8"),
}

# Schrift mit Umlauten und Gedankenstrich fürs PDF: Server (DejaVu), Mac (Arial). Ohne TTF fällt
# fpdf auf Helvetica zurück, die nur Latin-1 kann – dann werden Zeichen ersetzt.
_SCHRIFTEN = (
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
)


def _d(z: dict) -> date | None:
    try:
        return date.fromisoformat(z.get("datum", ""))
    except ValueError:
        return None


def datum_text(z: dict) -> str:
    d = _d(z)
    return f"{WT_KURZ[d.weekday()]} {d.strftime('%d.%m.%Y')}" if d else z.get("datum", "")


def zeit_text(z: dict) -> str:
    u, b = z.get("uhrzeit", ""), z.get("uhrzeit_bis", "")
    return f"{u}–{b}" if u and b else u


def _sortiert(zeilen: list[dict]) -> list[dict]:
    return sorted(zeilen, key=lambda z: (z.get("datum", ""), z.get("uhrzeit", ""), z.get("verein_name", "")))


def _nach_monat(zeilen: list[dict]):
    monat, gruppe = None, []
    for z in _sortiert(zeilen):
        d = _d(z)
        m = (d.year, d.month) if d else None
        if gruppe and m != monat:
            yield monat, gruppe
            gruppe = []
        monat = m
        gruppe.append(z)
    if gruppe:
        yield monat, gruppe


def _monat_text(m) -> str:
    return f"{MONATE[m[1] - 1]} {m[0]}" if m else "Ohne Datum"


def _spalten(mit_verein: bool) -> list[tuple[str, str]]:
    sp = [("Datum", "datum"), ("Uhrzeit", "zeit"), ("Termin", "bezeichnung"), ("Ort", "ort")]
    if mit_verein:
        sp.insert(3, ("Verein", "verein_name"))
    return sp


def _wert(z: dict, feld: str) -> str:
    if feld == "datum":
        return datum_text(z)
    if feld == "zeit":
        return zeit_text(z)
    return str(z.get(feld, "") or "")


# ── Formate ──────────────────────────────────────────────────────────────────

def _xlsx(titel: str, zeilen: list[dict], mit_verein: bool) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Termine"
    ws.append([titel])
    ws["A1"].font = Font(bold=True, size=14)
    sp = _spalten(mit_verein)
    ws.append([k for k, _ in sp])
    for zelle in ws[2]:
        zelle.font = Font(bold=True)
    for z in _sortiert(zeilen):
        ws.append([_wert(z, f) for _, f in sp])
    for i, breite in enumerate([16, 12, 40, 30, 30][:len(sp)], 1):
        ws.column_dimensions[chr(64 + i)].width = breite
    ws.freeze_panes = "A3"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _docx(titel: str, zeilen: list[dict], mit_verein: bool) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.add_heading(titel, level=1)
    sp = _spalten(mit_verein)
    for monat, gruppe in _nach_monat(zeilen):
        doc.add_heading(_monat_text(monat), level=2)
        tab = doc.add_table(rows=1, cols=len(sp))
        tab.style = "Light List Accent 1"
        for i, (kopf, _) in enumerate(sp):
            tab.rows[0].cells[i].text = kopf
        for z in gruppe:
            zellen = tab.add_row().cells
            for i, (_, f) in enumerate(sp):
                zellen[i].text = _wert(z, f)
        for zeile in tab.rows:
            for zelle in zeile.cells:
                for p in zelle.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(10)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _ods(titel: str, zeilen: list[dict], mit_verein: bool) -> bytes:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell
    from odf.text import P
    doc = OpenDocumentSpreadsheet()
    tab = Table(name="Termine")

    def zeile(werte):
        tr = TableRow()
        for w in werte:
            tc = TableCell(valuetype="string")
            tc.addElement(P(text=w))
            tr.addElement(tc)
        tab.addElement(tr)

    zeile([titel])
    sp = _spalten(mit_verein)
    zeile([k for k, _ in sp])
    for z in _sortiert(zeilen):
        zeile([_wert(z, f) for _, f in sp])
    doc.spreadsheet.addElement(tab)
    out = io.BytesIO()
    doc.write(out)
    return out.getvalue()


def _odt(titel: str, zeilen: list[dict], mit_verein: bool) -> bytes:
    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties
    from odf.text import H, P, Span
    doc = OpenDocumentText()
    fett = Style(name="Fett", family="text")
    fett.addElement(TextProperties(fontweight="bold"))
    doc.automaticstyles.addElement(fett)
    doc.text.addElement(H(outlinelevel=1, text=titel))
    for monat, gruppe in _nach_monat(zeilen):
        doc.text.addElement(H(outlinelevel=2, text=_monat_text(monat)))
        for z in gruppe:
            p = P()
            p.addElement(Span(stylename=fett, text=f"{datum_text(z)}  {zeit_text(z)}  "))
            rest = z.get("bezeichnung", "")
            if mit_verein and z.get("verein_name"):
                rest += f" – {z['verein_name']}"
            if z.get("ort"):
                rest += f" ({z['ort']})"
            p.addText(rest)
            doc.text.addElement(p)
    out = io.BytesIO()
    doc.write(out)
    return out.getvalue()


def _pdf(titel: str, zeilen: list[dict], mit_verein: bool) -> bytes:
    from fpdf import FPDF
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    schrift = "Helvetica"
    for normal, fett in _SCHRIFTEN:
        if Path(normal).exists() and Path(fett).exists():
            pdf.add_font("Text", "", normal)
            pdf.add_font("Text", "B", fett)
            schrift = "Text"
            break

    def txt(s: str) -> str:
        if schrift != "Helvetica":
            return s
        s = s.replace("–", "-").replace("„", '"').replace("“", '"').replace("…", "...")
        return s.encode("latin-1", "replace").decode("latin-1")

    pdf.set_title(titel)
    pdf.add_page()
    pdf.set_font(schrift, "B", 16)
    pdf.multi_cell(0, 8, txt(titel), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    sp = _spalten(mit_verein)
    breiten = {"datum": 28, "zeit": 22, "bezeichnung": 60, "verein_name": 40, "ort": 40} if mit_verein \
        else {"datum": 28, "zeit": 22, "bezeichnung": 80, "ort": 60}
    for monat, gruppe in _nach_monat(zeilen):
        pdf.ln(2)
        pdf.set_font(schrift, "B", 12)
        pdf.cell(0, 7, txt(_monat_text(monat)), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(schrift, "", 9)
        with pdf.table(col_widths=[breiten[f] for _, f in sp], text_align="LEFT",
                       line_height=5, first_row_as_headings=True) as tab:
            kopf = tab.row()
            for k, _ in sp:
                kopf.cell(txt(k))
            for z in gruppe:
                r = tab.row()
                for _, f in sp:
                    r.cell(txt(_wert(z, f)))
    pdf.set_y(-12)
    pdf.set_font(schrift, "", 7)
    pdf.cell(0, 5, txt(f"Erstellt mit vereinskalender.online · {date.today().strftime('%d.%m.%Y')}"), align="C")
    return bytes(pdf.output())


def _ics_escape(text: str) -> str:
    return (str(text).replace("\\", "\\\\").replace(",", "\\,").replace(";", "\\;")
            .replace("\r\n", "\\n").replace("\n", "\\n").replace("\r", "\\n"))


def _ics_zeiten(tag: date, uhrzeit: str, uhrzeit_bis: str = "") -> tuple[str, str]:
    def stamp(t: datetime) -> str:
        return t.strftime("%Y%m%dT%H%M00")
    if not uhrzeit:
        return (f"DTSTART;VALUE=DATE:{tag.strftime('%Y%m%d')}",
                f"DTEND;VALUE=DATE:{(tag + timedelta(days=1)).strftime('%Y%m%d')}")
    start = datetime.combine(tag, datetime.strptime(uhrzeit, "%H:%M").time())
    if uhrzeit_bis:
        ende = datetime.combine(tag, datetime.strptime(uhrzeit_bis, "%H:%M").time())
        if ende <= start:
            ende += timedelta(days=1)
    else:
        ende = min(start + timedelta(hours=1), datetime.combine(tag, datetime.max.time()).replace(second=0, microsecond=0))
    return f"DTSTART:{stamp(start)}", f"DTEND:{stamp(ende)}"


def _ics(titel: str, zeilen: list[dict], mit_verein: bool) -> bytes:
    jetzt = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    zeilen_ics = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Vereinskalender//Jahresplanung//DE",
                  "CALSCALE:GREGORIAN", f"X-WR-CALNAME:{_ics_escape(titel)}", "X-WR-TIMEZONE:Europe/Berlin"]
    vergeben = set()
    for z in _sortiert(zeilen):
        d = _d(z)
        if not d:
            continue
        try:
            start, ende = _ics_zeiten(d, z.get("uhrzeit", ""), z.get("uhrzeit_bis", ""))
        except ValueError:
            start, ende = _ics_zeiten(d, "", "")
        titel_z = z.get("bezeichnung") or "Termin"
        # UID nach demselben Schema wie der Abo-Feed (ADR-019), mit Endung bei Kollision
        uid = f"{z['datum']}-{re.sub(r'[^a-z0-9]', '', titel_z.lower()[:20])}-{z.get('verein', '')}"
        basis, n = uid, 2
        while uid in vergeben:
            uid, n = f"{basis}-{n}", n + 1
        vergeben.add(uid)
        beschr = z.get("verein_name", "") + (f"\n{z['ort']}" if z.get("ort") else "")
        zeilen_ics += ["BEGIN:VEVENT", f"UID:{uid}@vereinskalender", f"DTSTAMP:{jetzt}", start, ende,
                       f"SUMMARY:{_ics_escape(titel_z)}", f"DESCRIPTION:{_ics_escape(beschr)}"]
        if z.get("ort"):
            zeilen_ics.append(f"LOCATION:{_ics_escape(z['ort'])}")
        zeilen_ics.append("END:VEVENT")
    zeilen_ics.append("END:VCALENDAR")
    return "\r\n".join(zeilen_ics).encode("utf-8")


_SCHREIBER = {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "odt": _odt, "ods": _ods, "ics": _ics}


def exportiere(fmt: str, titel: str, zeilen: list[dict], mit_verein: bool = True) -> tuple[bytes, str, str]:
    """(Inhalt, Mimetype, Dateiendung). KeyError bei unbekanntem Format."""
    inhalt = _SCHREIBER[fmt](titel, zeilen, mit_verein)
    return inhalt, FORMATE[fmt][1], fmt


def dateiname(titel: str, fmt: str) -> str:
    return re.sub(r"[^\wäöüÄÖÜß.-]+", "_", titel).strip("_")[:80] + "." + fmt


# ── Ergebnis einer Planungsrunde (Besprechungsergebnis) ──────────────────────

STATUS_TEXT = {"entwurf": "Entwurf", "bestaetigt": "Bestätigt", "veroeffentlicht": "Veröffentlicht"}


def ergebnis_pdf(daten: dict) -> bytes:
    """Besprechungsergebnis aus dem eingefrorenen Stand einer Runde (`runde_ergebnis.daten_json`).
    Kopf (Runde, Organisator, Abschluss, Teilnehmer), Termine nach Monat mit Status, offene Konflikte, Verlauf.
    Rein aus den Daten erzeugt – derselbe Stand ergibt immer dasselbe Dokument."""
    from fpdf import FPDF
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    schrift = "Helvetica"
    for normal, fett in _SCHRIFTEN:
        if Path(normal).exists() and Path(fett).exists():
            pdf.add_font("Text", "", normal)
            pdf.add_font("Text", "B", fett)
            schrift = "Text"
            break

    def txt(s) -> str:
        s = str(s or "")
        if schrift != "Helvetica":
            return s
        s = s.replace("–", "-").replace("→", "->").replace("„", '"').replace("“", '"').replace("…", "...")
        return s.encode("latin-1", "replace").decode("latin-1")

    def zeitpunkt(iso: str) -> str:
        try:
            return datetime.fromisoformat(iso).strftime("%d.%m.%Y, %H:%M Uhr")
        except ValueError:
            return iso

    titel = f"Ergebnis: {daten['name']} {daten['jahr']}"
    pdf.set_title(titel)
    pdf.set_creation_date(datetime.fromisoformat(daten["abgeschlossen_am"]))   # gleicher Stand = gleiche Datei
    pdf.add_page()
    pdf.set_font(schrift, "B", 16)
    pdf.multi_cell(0, 8, txt(titel), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(schrift, "", 10)
    kopf = [("Organisator", daten["organisator_name"]),
            ("Abgeschlossen", f"{zeitpunkt(daten['abgeschlossen_am'])} von {daten['abgeschlossen_von_name']}"
                              f" (Version {daten.get('version', 1)})"),
            ("Teilnehmer", ", ".join(t["name"] for t in daten["teilnehmer"]))]
    for k, w in kopf:
        pdf.set_font(schrift, "B", 10)
        pdf.cell(32, 6, txt(k))
        pdf.set_font(schrift, "", 10)
        pdf.multi_cell(0, 6, txt(w), new_x="LMARGIN", new_y="NEXT")
    termine = daten["termine"]
    n = {s: sum(1 for t in termine if t["status"] == s) for s in STATUS_TEXT}
    pdf.ln(2)
    pdf.set_font(schrift, "", 9)
    pdf.multi_cell(0, 5, txt(f"{len(termine)} Termine: {n['veroeffentlicht']} veröffentlicht, {n['bestaetigt']} bestätigt, "
                             f"{n['entwurf']} noch Entwurf. Bestätigt = vom Verein als vereinbart gemeldet; "
                             f"veröffentlicht = steht im Vereinskalender."), new_x="LMARGIN", new_y="NEXT")

    breiten = [26, 18, 62, 44, 24]
    for monat, gruppe in _nach_monat(termine):
        pdf.ln(2)
        pdf.set_font(schrift, "B", 12)
        pdf.cell(0, 7, txt(_monat_text(monat)), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(schrift, "", 9)
        with pdf.table(col_widths=breiten, text_align="LEFT", line_height=5, first_row_as_headings=True) as tab:
            k = tab.row()
            for h in ("Datum", "Uhrzeit", "Termin", "Verein", "Status"):
                k.cell(txt(h))
            for t in gruppe:
                r = tab.row()
                termin = t.get("bezeichnung", "") + (f" ({t['ort']})" if t.get("ort") else "")
                for w in (datum_text(t), zeit_text(t), termin, t.get("verein_name", ""), STATUS_TEXT.get(t["status"], "")):
                    r.cell(txt(w))

    pdf.ln(3)
    pdf.set_font(schrift, "B", 12)
    pdf.cell(0, 7, txt(f"Offene Konflikte am gleichen Tag ({len(daten['konflikte'])})"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(schrift, "", 9)
    if not daten["konflikte"]:
        pdf.cell(0, 5, txt("Keine."), new_x="LMARGIN", new_y="NEXT")
    for k in daten["konflikte"]:
        pdf.multi_cell(0, 5, txt(f"{datum_text(k)}: {k['a']} – {k['b']}"), new_x="LMARGIN", new_y="NEXT")

    pdf.ln(3)
    pdf.set_font(schrift, "B", 12)
    pdf.cell(0, 7, txt("Verlauf der Runde"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(schrift, "", 8)
    with pdf.table(col_widths=[34, 46, 94], text_align="LEFT", line_height=4.5, first_row_as_headings=True) as tab:
        k = tab.row()
        for h in ("Zeit", "Verein", "Was"):
            k.cell(txt(h))
        for p in daten["verlauf"]:
            r = tab.row()
            for w in (zeitpunkt(p["zeit"]), p["verein_name"], p["aktion"] + (f": {p['details']}" if p["details"] else "")):
                r.cell(txt(w))
    pdf.set_y(-12)
    pdf.set_font(schrift, "", 7)
    pdf.cell(0, 5, txt(f"Vereinskalender.online · Planungsrunde „{daten['name']}“ · Version {daten.get('version', 1)}"), align="C")
    return bytes(pdf.output())
